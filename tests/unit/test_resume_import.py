"""Resume text extraction from the four accepted formats."""

from __future__ import annotations

import io
import zipfile

import pytest

from shortlist.resume_import import (
    MAX_RESUME_BYTES,
    MAX_RESUME_TEXT_CHARS,
    MAX_RESUME_UNCOMPRESSED_BYTES,
    UnsupportedResumeError,
    extract_text,
)


def test_markdown_passes_through() -> None:
    assert extract_text(b"# David\n\nEngineer", "resume.md") == "# David\n\nEngineer"


def test_plain_text_passes_through() -> None:
    assert extract_text(b"David, engineer", "resume.txt") == "David, engineer"


def test_pdf_is_extracted() -> None:
    pytest.importorskip("reportlab")
    from reportlab.pdfgen.canvas import Canvas

    buf = io.BytesIO()
    canvas = Canvas(buf)
    canvas.drawString(72, 720, "David Miller Senior Engineer")
    canvas.save()
    assert "David Miller" in extract_text(buf.getvalue(), "resume.pdf")


def test_docx_is_extracted() -> None:
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("David Miller")
    doc.add_paragraph("Senior Engineer")
    doc.save(buf)
    text = extract_text(buf.getvalue(), "resume.docx")
    assert "David Miller" in text
    assert "Senior Engineer" in text


def test_unknown_extension_rejected() -> None:
    with pytest.raises(UnsupportedResumeError):
        extract_text(b"anything", "resume.rtf")


def test_renamed_executable_rejected_by_content() -> None:
    """An ELF binary renamed to .pdf must not be accepted on extension alone."""
    with pytest.raises(UnsupportedResumeError):
        extract_text(b"\x7fELF\x02\x01\x01\x00" + b"\x00" * 64, "resume.pdf")


def test_oversize_rejected() -> None:
    with pytest.raises(UnsupportedResumeError):
        extract_text(b"x" * (MAX_RESUME_BYTES + 1), "resume.md")


def test_undecodable_text_rejected() -> None:
    with pytest.raises(UnsupportedResumeError):
        extract_text(b"\xff\xfe\x00\x00binary", "resume.txt")


def test_corrupt_pdf_rejected() -> None:
    """Right magic bytes, broken internals: pypdf's own error must not escape."""
    pytest.importorskip("reportlab")
    from reportlab.pdfgen.canvas import Canvas

    buf = io.BytesIO()
    canvas = Canvas(buf)
    canvas.drawString(72, 720, "David Miller Senior Engineer")
    canvas.save()
    data = buf.getvalue()
    truncated = data[: len(data) // 2]
    with pytest.raises(UnsupportedResumeError):
        extract_text(truncated, "resume.pdf")


def test_corrupt_docx_rejected() -> None:
    """Right magic bytes, broken zip/xml internals: the parser's own error must not escape."""
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("David Miller")
    doc.save(buf)
    data = buf.getvalue()
    truncated = data[: len(data) // 2]
    with pytest.raises(UnsupportedResumeError):
        extract_text(truncated, "resume.docx")


def _compressible_docx(n_paragraphs: int, chunk_len: int) -> bytes:
    """A .docx of `n_paragraphs` paragraphs of a repeated character, highly
    compressible so the zip stays tiny while the declared uncompressed size
    and the extractable text both scale with `n_paragraphs * chunk_len`.
    """
    from docx import Document

    buf = io.BytesIO()
    doc = Document()
    chunk = "x" * chunk_len
    for _ in range(n_paragraphs):
        doc.add_paragraph(chunk)
    doc.save(buf)
    return buf.getvalue()


def test_oversized_uncompressed_rejected_before_parsing() -> None:
    """A docx whose declared uncompressed size exceeds the cap must be rejected
    before `Document(...)` ever parses it - the whole point of checking the
    zip's central directory first.
    """
    data = _compressible_docx(n_paragraphs=40, chunk_len=1_000_000)
    assert len(data) < MAX_RESUME_BYTES  # a compression attack, not just a big file
    with pytest.raises(UnsupportedResumeError, match="uncompressed content"):
        extract_text(data, "resume.docx")


def test_oversized_extraction_rejected() -> None:
    """A docx under the uncompressed-size cap but whose extracted text still
    exceeds MAX_RESUME_TEXT_CHARS must be caught by the extraction-time cap -
    proving that bound stays reachable once the pre-parse one exists.
    """
    data = _compressible_docx(n_paragraphs=20, chunk_len=100_000)
    assert len(data) < MAX_RESUME_BYTES
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        uncompressed = sum(info.file_size for info in zf.infolist())
    assert uncompressed < MAX_RESUME_UNCOMPRESSED_BYTES
    assert MAX_RESUME_TEXT_CHARS < 20 * 100_000
    with pytest.raises(UnsupportedResumeError, match="extracted text"):
        extract_text(data, "resume.docx")
