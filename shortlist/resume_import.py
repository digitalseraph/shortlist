"""Turn an uploaded resume into the plain text the ranking pipeline reads.

`embed_resume_cached` calls `read_text()` on the stored file, so every accepted
format is normalised to text here. PDF extraction is lossy - the caller is
expected to show the result for review before saving it.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterable
from pathlib import Path

MAX_RESUME_BYTES = 2 * 1024 * 1024
# A resume runs a few thousand characters. This is generous headroom while
# still refusing a small, highly compressible file (a zip/docx "bomb") that
# would otherwise expand into megabytes of text during extraction.
MAX_RESUME_TEXT_CHARS = 1_000_000
# python-docx parses a .docx's whole XML tree into memory before yielding a
# single paragraph, so MAX_RESUME_TEXT_CHARS alone is too late: a small zip
# can already have cost hundreds of MB by the time the first paragraph is
# checked. Reading the zip's central directory (no decompression) catches
# that before `Document(...)` is ever called.
MAX_RESUME_UNCOMPRESSED_BYTES = 32 * 1024 * 1024

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"  # a .docx is a zip container


class UnsupportedResumeError(ValueError):
    """The upload is not a resume we can read."""


def _as_text(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise UnsupportedResumeError("file is not valid UTF-8 text") from exc


def _join_capped(parts: Iterable[str], sep: str) -> str:
    """Join `parts`, raising as soon as the running total exceeds the cap.

    Checks the running total piece by piece instead of joining first and
    measuring after, so a file that decompresses into something huge is
    rejected before that huge string is ever fully materialised.
    """
    out: list[str] = []
    total = 0
    for part in parts:
        total += len(part)
        if total > MAX_RESUME_TEXT_CHARS:
            raise UnsupportedResumeError(
                f"extracted text is larger than {MAX_RESUME_TEXT_CHARS:,} characters"
            )
        out.append(part)
    return sep.join(out).strip()


def _from_pdf(data: bytes) -> str:
    if not data.startswith(_PDF_MAGIC):
        raise UnsupportedResumeError("file does not look like a PDF")
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        return _join_capped((page.extract_text() or "" for page in reader.pages), "\n\n")
    except UnsupportedResumeError:
        raise
    except Exception as exc:
        raise UnsupportedResumeError("file could not be read as a PDF") from exc


def _from_docx(data: bytes) -> str:
    if not data.startswith(_ZIP_MAGIC):
        raise UnsupportedResumeError("file does not look like a .docx")
    from docx import Document

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            uncompressed = sum(info.file_size for info in zf.infolist())
        if uncompressed > MAX_RESUME_UNCOMPRESSED_BYTES:
            raise UnsupportedResumeError(
                "uncompressed content is larger than "
                f"{MAX_RESUME_UNCOMPRESSED_BYTES // (1024 * 1024)} MB"
            )
        doc = Document(io.BytesIO(data))
        return _join_capped((p.text for p in doc.paragraphs), "\n")
    except UnsupportedResumeError:
        raise
    except Exception as exc:
        raise UnsupportedResumeError("file could not be read as a .docx") from exc


def extract_text(data: bytes, filename: str) -> str:
    """Extract resume text, validating by extension and by leading bytes."""
    if len(data) > MAX_RESUME_BYTES:
        raise UnsupportedResumeError(f"file is larger than {MAX_RESUME_BYTES // (1024 * 1024)} MB")
    suffix = Path(filename).suffix.lower()
    if suffix in (".md", ".txt"):
        return _as_text(data)
    if suffix == ".pdf":
        return _from_pdf(data)
    if suffix == ".docx":
        return _from_docx(data)
    raise UnsupportedResumeError(f"unsupported file type {suffix or filename!r}")
